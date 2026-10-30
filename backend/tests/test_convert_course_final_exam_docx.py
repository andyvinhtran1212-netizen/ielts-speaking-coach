"""Parsing tests for the Course final .docx → assessment JSON converter.

The fixtures rebuild the real paper/key layout (Heading 1 parts followed by a
range line, Heading 2 passages, bold ``Câu NNN`` / ``A.`` labels, a bold
``Ngữ cảnh`` passage label, the quick-key table and the detailed key with
``Nguồn x.y.zz``) without committing the exam content itself.
"""

from __future__ import annotations

import copy

import docx
import pytest

from scripts.convert_course_final_exam_docx import build_rows, read_docx
from services.course_assessment_import import normalize_assessment_rows

PASSAGES = [
    # (passage lines as (bold label, text)), questions as (stem, options, key)
    ([("", "We travel in July.")], [
        ("In July cung cấp loại thông tin nào?",
         ["Phương tiện.", "Người đi cùng.", "Tháng đi.", "Giờ đi.", "Nơi đến."], "C"),
        ("Đổi July thành Monday thì dùng giới từ nào?",
         ["On Monday.", "Into Monday.", "By Monday.", "Under Monday.", "At Monday."], "A"),
    ]),
    ([("Ngữ cảnh", "We visited a museum."), ("", "The museum was quiet.")], [
        ("The museum chỉ bảo tàng nào?",
         ["Bảo tàng mai mở.", "Bảo tàng vừa nhắc.", "Mọi bảo tàng.",
          "Bảo tàng khác.", "Không bảo tàng nào."], "B"),
        ("The dùng vì sao?",
         ["Số nhiều.", "Tên riêng.", "Đã xác định.", "Không đếm được.", "Lần đầu nhắc."],
         "C"),
    ]),
]
SOURCES = ["6.2.02", "17.4.10"]


def _para(document, labelled, style=None):
    paragraph = document.add_paragraph(style=style)
    label, text = labelled
    if label:
        paragraph.add_run(f"{label} ").bold = True
    paragraph.add_run(text)


def _write_paper(path, passages=PASSAGES):
    document = docx.Document()
    document.add_paragraph("Kiểm tra tổng kết", style="Title")
    _para(document, ("", "Chọn một đáp án phù hợp nhất A, B, C, D hoặc E."))
    number = 0
    for p_index, (lines, questions) in enumerate(passages, start=1):
        if p_index == 1:
            document.add_paragraph("Phần 1", style="Heading 1")
            _para(document, ("", "Ngữ liệu 001 đến 001"))
        if p_index == 2:
            document.add_paragraph("Phần 2", style="Heading 1")
            _para(document, ("", "Ngữ liệu 002 đến 002"))
        document.add_paragraph(f"Ngữ liệu {p_index:03d}", style="Heading 2")
        for label, text in lines:
            paragraph = document.add_paragraph()
            if label:
                paragraph.add_run(f"{label} ").bold = True
            paragraph.add_run(text).italic = True
        for stem, options, _key in questions:
            number += 1
            _para(document, (f"Câu {number:03d}", stem))
            for letter, option in zip("ABCDE", options):
                _para(document, (f"{letter}.", option))
    document.save(path)


def _write_key(path, passages=PASSAGES, quick_override=None, detail_override=None):
    quick_override = quick_override or {}
    detail_override = detail_override or {}
    document = docx.Document()
    document.add_paragraph("Đáp án và giải thích", style="Title")
    document.add_paragraph("Bảng đáp án nhanh", style="Heading 1")
    answers = [key for _lines, questions in passages for *_rest, key in questions]
    table = document.add_table(rows=1 + len(answers), cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Câu", "Đáp án"
    for row, key in enumerate(answers, start=1):
        table.cell(row, 0).text = f"{row:03d}"
        table.cell(row, 1).text = quick_override.get(row, key)
    document.add_paragraph("Đáp án chi tiết theo ngữ liệu", style="Heading 1")
    number = 0
    for p_index, (_lines, questions) in enumerate(passages, start=1):
        document.add_paragraph(
            f"Ngữ liệu {p_index:03d}  Nguồn {SOURCES[p_index - 1]}", style="Heading 2")
        _para(document, ("Dịch tham khảo", f"Bản dịch {p_index}."))
        for _stem, options, key in questions:
            number += 1
            letter, text = detail_override.get(number, (key, options["ABCDE".index(key)]))
            _para(document, (f"Câu {number:03d}  {letter}  Đáp án", text))
            _para(document, ("Giải thích", f"Lý do câu {number}."))
        _para(document, ("Chấp nhận thêm khi tự luận", "Chỉ cho bản tự luận."))
    document.save(path)


def _build(tmp_path, *, passages=PASSAGES, key_kwargs=None, **options):
    paper, key = tmp_path / "a.docx", tmp_path / "b.docx"
    _write_paper(paper, passages)
    _write_key(key, passages, **(key_kwargs or {}))
    paper_paras, paper_tables = read_docx(paper)
    key_paras, key_tables = read_docx(key)
    return build_rows(paper_paras, paper_tables, key_paras, key_tables,
                      id_prefix="C1-FINAL", expected_questions=4,
                      expected_passages=2, **options)


def test_maps_the_real_layout_into_rows_with_passage_key_and_explanation(tmp_path):
    rows, errors, report = _build(tmp_path)
    assert errors == []
    assert [row["id"] for row in rows] == [f"C1-FINAL-Q00{n}" for n in range(1, 5)]
    first, third = rows[0], rows[2]
    assert first["de"] == ("Ngữ liệu 001\nWe travel in July.\n\n"
                           "In July cung cấp loại thông tin nào?")
    assert third["de"].startswith(
        "Ngữ liệu 002\nNgữ cảnh: We visited a museum.\nThe museum was quiet.\n\n")
    assert first["dap_an"] == 2 and first["dap_an_chu"] == "C"
    assert first["pa"][2] == "Tháng đi."
    assert first["giai_thich"] == "Lý do câu 1.\nDịch tham khảo: Bản dịch 1."
    assert "Chỉ cho bản tự luận" not in rows[1]["giai_thich"]
    assert (first["lesson_primary"], third["lesson_primary"]) == ("B06", "B17")
    assert (third["truc"], third["diem_day"]) == ("Buổi 17", "Nguồn 17.4.10")
    assert [row["phan"] for row in rows] == [1, 1, 2, 2]
    assert report["by_part"] == {1: 2, 2: 2}


def test_flags_quick_key_disagreeing_with_the_detailed_key(tmp_path):
    rows, errors, _ = _build(tmp_path, key_kwargs={"quick_override": {2: "E"}})
    assert errors == ["câu 002: bảng nhanh E ≠ đáp án chi tiết A"]
    assert [row["id"] for row in rows] == ["C1-FINAL-Q001", "C1-FINAL-Q003", "C1-FINAL-Q004"]


def test_flags_answer_text_that_is_not_the_keyed_option(tmp_path):
    _, errors, _ = _build(tmp_path, key_kwargs={"detail_override": {3: ("B", "Mọi bảo tàng.")}})
    assert len(errors) == 1 and errors[0].startswith("câu 003: phương án B trong đề")


def test_flags_a_question_without_five_options_and_a_numbering_gap(tmp_path):
    broken = copy.deepcopy(PASSAGES)
    stem, options, key = broken[1][1][1]
    broken[1][1][1] = (stem, options[:4], key)
    _, errors, _ = _build(tmp_path, passages=broken)
    assert "câu 004 có 4 phương án, cần 5" in errors

    _, errors, _ = _build(tmp_path, passages=PASSAGES)
    assert errors == []
    paper_paras, paper_tables = read_docx(tmp_path / "a.docx")
    key_paras, key_tables = read_docx(tmp_path / "b.docx")
    _, errors, _ = build_rows(paper_paras, paper_tables, key_paras, key_tables,
                              id_prefix="C1-FINAL", expected_questions=5,
                              expected_passages=2)
    assert "đề: câu thiếu 005" in errors


def test_default_output_is_refused_by_the_importer_until_reviewed_fields_exist(tmp_path):
    rows, errors, _ = _build(tmp_path)
    assert errors == []
    assert not {"dang", "bay", "ma_bay", "classroom_alignment"} & set(rows[0])
    with pytest.raises(ValueError, match="thiếu mã dạng"):
        normalize_assessment_rows(rows, expected_count=4, scope="B01-B20")

    rows, errors, _ = _build(tmp_path, dang="A1", generic_wrong_reason=True,
                             confirm_alignment=True)
    assert errors == []
    mapped = normalize_assessment_rows(rows, expected_count=4, scope="B01-B20")
    assert set(mapped[0]["why_wrong"]) == {"0", "1", "3", "4"}
    assert rows[0]["ma_bay"][2] == "" and len(set(rows[0]["ma_bay"])) == 5
