"""Convert the Course final exam (.docx question paper + .docx key) into the
five-option assessment JSON read by ``import_course_assessment_bank``.

The question paper ("đề A") has ``Ngữ liệu NNN`` headings, one italic passage
(optionally with a bold ``Ngữ cảnh`` / ``Câu kế tiếp`` label), then two
``Câu NNN`` questions with options ``A.``–``E.``.  The key ("đề B") has a quick
answer table plus, per passage, ``Nguồn x.y.zz``, a reference translation and
``Câu NNN  X  Đáp án …`` / ``Giải thích …`` lines.

Every cross-check failure is collected and reported; nothing is guessed::

    python -m scripts.convert_course_final_exam_docx \
      --paper de-A.docx --key de-B.docx --out FINAL.json --id-prefix C1-FINAL

The source has no per-distractor reasons, no question-type code and no
classroom-alignment sign-off, which the assessment contract requires.  Those
fields are omitted unless the operator passes ``--generic-wrong-reason``,
``--dang`` and ``--confirm-classroom-alignment`` explicitly, so the importer's
dry-run refuses an unreviewed file instead of importing invented content.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

LETTERS = "ABCDE"
_PASSAGE_HEAD = re.compile(r"Ngữ liệu (\d{3})")
_KEY_PASSAGE_HEAD = re.compile(r"Ngữ liệu (\d{3})\s+Nguồn (\d+)\.(\d+)\.(\d+)")
_QUESTION = re.compile(r"Câu (\d{3}) (.+)", re.S)
_OPTION = re.compile(r"([A-E])\. (.+)", re.S)
_KEY_ANSWER = re.compile(r"Câu (\d{3})\s+([A-E])\s+Đáp án (.+)", re.S)
_PART_HEAD = re.compile(r"Phần (\d)")
# Generic, truthful wrong-answer note used only when the operator opts in.
GENERIC_WRONG_REASON = (
    "Phương án này không khớp dữ kiện hoặc cấu trúc của câu nguồn; "
    "đối chiếu với phần giải thích đáp án đúng."
)


@dataclass
class Para:
    style: str
    text: str
    label: str = ""          # leading bold run text, e.g. "Câu 001", "Ngữ cảnh"


@dataclass
class Question:
    number: int
    passage: int
    part: int
    stem: str
    options: list[str] = field(default_factory=list)


@dataclass
class Passage:
    number: int
    part: int
    lines: list[str] = field(default_factory=list)


def _norm(text: str) -> str:
    return " ".join(text.replace(" ", " ").split())


def read_docx(path: Path) -> tuple[list[Para], list[list[str]]]:
    """Paragraphs (with their leading bold label) and table cell texts."""
    import docx  # python-docx, already a backend dependency

    document = docx.Document(str(path))
    paras: list[Para] = []
    for paragraph in document.paragraphs:
        text = _norm(paragraph.text)
        if not text:
            continue
        label = ""
        for run in paragraph.runs:
            if not run.text.strip() and not label:
                continue
            if run.bold:
                label += run.text
            else:
                break
        paras.append(Para(paragraph.style.name, text, _norm(label)))
    tables = [[_norm(cell.text) for row in table.rows for cell in row.cells]
              for table in document.tables]
    return paras, tables


def parse_paper(paras: list[Para], errors: list[str]):
    passages: dict[int, Passage] = {}
    questions: dict[int, Question] = {}
    part = 0
    current_passage: Passage | None = None
    current_question: Question | None = None
    for para in paras:
        if para.style.startswith("Heading 1"):
            match = _PART_HEAD.fullmatch(para.text)
            if match:
                part = int(match.group(1))
            # A part heading is followed by its range line, not by passage text.
            current_passage = current_question = None
            continue
        head = _PASSAGE_HEAD.fullmatch(para.text)
        if para.style.startswith("Heading 2") and head:
            number = int(head.group(1))
            current_passage = Passage(number, part)
            if number in passages:
                # Keep the first block; the duplicate's lines go to an unsaved Passage.
                errors.append(f"đề: ngữ liệu {number:03d} xuất hiện hai lần")
            else:
                passages[number] = current_passage
            current_question = None
            continue
        if current_passage is None:
            continue  # cover page, instructions and part table
        question = _QUESTION.fullmatch(para.text)
        if question and para.label.startswith("Câu "):
            number = int(question.group(1))
            current_question = Question(
                number, current_passage.number, current_passage.part,
                _norm(question.group(2)))
            if number in questions:
                # Keep the first question; the duplicate's options go to an unsaved Question.
                errors.append(f"đề: câu {number:03d} xuất hiện hai lần")
            else:
                questions[number] = current_question
            continue
        option = _OPTION.fullmatch(para.text)
        if option and para.label == f"{option.group(1)}.":
            if current_question is None:
                errors.append(f"đề: phương án ngoài câu hỏi ở ngữ liệu "
                              f"{current_passage.number:03d}")
                continue
            expected = LETTERS[len(current_question.options)] \
                if len(current_question.options) < 5 else "?"
            if option.group(1) != expected:
                errors.append(f"đề: câu {current_question.number:03d} có phương án "
                              f"{option.group(1)} ở vị trí của {expected}")
            current_question.options.append(_norm(option.group(2)))
            continue
        if current_question is not None:
            errors.append(f"đề: câu {current_question.number:03d} có dòng lạ: "
                          f"{para.text[:60]!r}")
            continue
        line = para.text
        if para.label and line.startswith(para.label):
            line = f"{para.label}: {_norm(line[len(para.label):])}"
        current_passage.lines.append(line)
    return passages, questions


def parse_quick_key(tables: list[list[str]], errors: list[str]) -> dict[int, str]:
    """``Câu | Đáp án`` pairs from the 'Bảng đáp án nhanh' table."""
    key: dict[int, str] = {}
    for cells in tables:
        if "Đáp án" not in cells:
            continue
        for number, letter in zip(cells[::2], cells[1::2]):
            if number == "Câu" and letter == "Đáp án":
                continue
            if not re.fullmatch(r"\d{3}", number) or letter not in LETTERS:
                errors.append(f"bảng nhanh: ô lạ {number!r} → {letter!r}")
                continue
            if int(number) in key:
                errors.append(f"bảng nhanh: câu {number} xuất hiện hai lần")
            key[int(number)] = letter
    return key


def parse_detailed_key(paras: list[Para], errors: list[str]):
    """Per-question letter, answer text, explanation; per-passage source."""
    sources: dict[int, tuple[int, int, int]] = {}
    translations: dict[int, str] = {}
    answers: dict[int, dict] = {}
    started = False
    passage = None
    last: dict | None = None
    for para in paras:
        if para.style.startswith("Heading 1"):
            started = para.text.startswith("Đáp án chi tiết")
            continue
        if not started:
            continue
        head = _KEY_PASSAGE_HEAD.fullmatch(para.text)
        if para.style.startswith("Heading 2"):
            if not head:
                errors.append(f"đáp án: tiêu đề ngữ liệu lạ {para.text!r}")
                passage = None
                continue
            passage = int(head.group(1))
            last = None
            if passage in sources:
                errors.append(f"đáp án: ngữ liệu {passage:03d} xuất hiện hai lần")
                passage = None
                continue
            sources[passage] = tuple(int(head.group(i)) for i in (2, 3, 4))
            continue
        if passage is None:
            continue
        if para.label == "Dịch tham khảo":
            translation = _norm(para.text[len(para.label):])
            if not translation:
                errors.append(f"đáp án: ngữ liệu {passage:03d} thiếu Dịch tham khảo")
                continue
            translations[passage] = translation
            continue
        answer = _KEY_ANSWER.fullmatch(para.text)
        if answer and para.label.startswith("Câu "):
            number = int(answer.group(1))
            if number in answers:
                errors.append(f"đáp án: câu {number:03d} xuất hiện hai lần")
            last = answers[number] = {
                "passage": passage, "letter": answer.group(2),
                "answer_text": _norm(answer.group(3)), "explanation": "",
            }
            continue
        if para.label == "Giải thích" and last is not None:
            if last["explanation"]:
                errors.append(f"đáp án: câu có hai dòng Giải thích ở ngữ liệu {passage:03d}")
            last["explanation"] = _norm(para.text[len(para.label):])
            continue
        if para.label == "Chấp nhận thêm khi tự luận":
            continue  # essay-version marking note, not part of the MCQ paper
        errors.append(f"đáp án: dòng lạ ở ngữ liệu {passage:03d}: {para.text[:60]!r}")
    return sources, translations, answers


def build_rows(paper_paras, paper_tables, key_paras, key_tables, *,
               id_prefix: str, expected_questions: int = 300,
               expected_passages: int = 150, dang: str | None = None,
               generic_wrong_reason: bool = False,
               confirm_alignment: bool = False):
    """Return ``(rows, errors, report)``; rows are only trusted if errors == []."""
    errors: list[str] = []
    passages, questions = parse_paper(paper_paras, errors)
    quick = parse_quick_key(key_tables, errors)
    sources, translations, detailed = parse_detailed_key(key_paras, errors)

    want_q = set(range(1, expected_questions + 1))
    want_p = set(range(1, expected_passages + 1))
    for label, found, want in (
        ("đề: câu", set(questions), want_q), ("đề: ngữ liệu", set(passages), want_p),
        ("bảng nhanh: câu", set(quick), want_q), ("đáp án: câu", set(detailed), want_q),
        ("đáp án: nguồn ngữ liệu", set(sources), want_p),
        ("đáp án: dịch ngữ liệu", set(translations), want_p),
    ):
        if missing := sorted(want - found):
            errors.append(f"{label} thiếu {', '.join(f'{n:03d}' for n in missing)}")
        if extra := sorted(found - want):
            errors.append(f"{label} thừa {', '.join(f'{n:03d}' for n in extra)}")

    rows: list[dict] = []
    for number in sorted(questions):
        question = questions[number]
        passage = passages[question.passage]
        if number != question.passage * 2 - 1 and number != question.passage * 2:
            errors.append(f"câu {number:03d} nằm ở ngữ liệu {question.passage:03d}, "
                          "không theo thứ tự 2 câu/ngữ liệu")
        if len(question.options) != 5:
            errors.append(f"câu {number:03d} có {len(question.options)} phương án, cần 5")
            continue
        if not passage.lines:
            errors.append(f"ngữ liệu {passage.number:03d} không có câu nguồn")
        quick_letter = quick.get(number)
        detail = detailed.get(number)
        if detail is None or quick_letter is None:
            continue
        if detail["letter"] != quick_letter:
            errors.append(f"câu {number:03d}: bảng nhanh {quick_letter} ≠ "
                          f"đáp án chi tiết {detail['letter']}")
            continue
        if detail["passage"] != question.passage:
            errors.append(f"câu {number:03d}: đề ở ngữ liệu {question.passage:03d}, "
                          f"đáp án ở ngữ liệu {detail['passage']:03d}")
        index = LETTERS.index(quick_letter)
        if _norm(question.options[index]).casefold() != detail["answer_text"].casefold():
            errors.append(f"câu {number:03d}: phương án {quick_letter} trong đề "
                          f"{question.options[index]!r} ≠ 'Đáp án' {detail['answer_text']!r}")
        if not detail["explanation"]:
            errors.append(f"câu {number:03d} thiếu Giải thích")
        lesson, unit, item = sources.get(question.passage, (0, 0, 0))
        explanation = detail["explanation"]
        if translations.get(question.passage):
            explanation += f"\nDịch tham khảo: {translations[question.passage]}"
        row = {
            "id": f"{id_prefix}-Q{number:03d}",
            "de": "\n".join([f"Ngữ liệu {passage.number:03d}", *passage.lines,
                             "", question.stem]),
            "pa": list(question.options),
            "dap_an": index,
            "dap_an_chu": quick_letter,
            "giai_thich": explanation,
            "muc": 1,
            "loai_cau_hoi": "multiple_choice_5",
            "so_phuong_an": 5,
            # "Nguồn x.y.zz": x is the Course lesson (1–20 in this paper).
            "lesson_primary": f"B{lesson:02d}",
            "diem_day": f"Nguồn {lesson}.{unit}.{item:02d}",
            # truc becomes item_key, which the Course report shows as the
            # grouping label, so group by lesson rather than 150 source codes.
            "truc": f"Buổi {lesson:02d}",
            "truc_nhom": f"B{lesson:02d}",
            "chieu": None,
            "phan": question.part,
            "ngu_lieu": passage.number,
        }
        if dang:
            row["dang"] = dang
        if generic_wrong_reason:
            row["bay"] = ["" if i == index else GENERIC_WRONG_REASON for i in range(5)]
            wrong = iter(range(1, 5))
            row["ma_bay"] = ["" if i == index else f"PA-SAI-{next(wrong)}"
                             for i in range(5)]
        if confirm_alignment:
            row["classroom_alignment"] = True
        rows.append(row)

    report = {
        "questions": len(questions), "passages": len(passages),
        "quick_key": len(quick), "detailed_key": len(detailed),
        "by_part": {p: sum(1 for r in rows if r["phan"] == p)
                    for p in sorted({r["phan"] for r in rows})},
        "by_lesson": {lesson: sum(1 for r in rows if r["lesson_primary"] == lesson)
                      for lesson in sorted({r["lesson_primary"] for r in rows})},
        "answer_letters": {letter: sum(1 for r in rows if r["dap_an_chu"] == letter)
                           for letter in LETTERS},
    }
    return rows, errors, report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--paper", required=True, help="đề trắc nghiệm .docx")
    parser.add_argument("--key", required=True, help="đáp án + giải thích .docx")
    parser.add_argument("--out", required=True)
    parser.add_argument("--id-prefix", required=True)
    parser.add_argument("--expected-count", type=int, default=300)
    parser.add_argument("--expected-passages", type=int, default=150)
    parser.add_argument("--dang", help="mã dạng gán cho mọi câu (chỉ khi đã duyệt)")
    parser.add_argument("--generic-wrong-reason", action="store_true",
                        help="điền lý do sai chung cho 4 phương án nhiễu (chỉ khi đã duyệt)")
    parser.add_argument("--confirm-classroom-alignment", action="store_true")
    args = parser.parse_args()

    paper_paras, paper_tables = read_docx(Path(args.paper))
    key_paras, key_tables = read_docx(Path(args.key))
    rows, errors, report = build_rows(
        paper_paras, paper_tables, key_paras, key_tables,
        id_prefix=args.id_prefix, expected_questions=args.expected_count,
        expected_passages=args.expected_passages, dang=args.dang,
        generic_wrong_reason=args.generic_wrong_reason,
        confirm_alignment=args.confirm_classroom_alignment,
    )
    print(json.dumps(report, ensure_ascii=False, indent=1))
    if errors:
        print(f"CROSS-CHECK FAIL: {len(errors)} lỗi", file=sys.stderr)
        for error in errors:
            print(f"  - {error}", file=sys.stderr)
        return 1
    Path(args.out).write_text(json.dumps(rows, ensure_ascii=False, indent=1) + "\n",
                              encoding="utf-8")
    print(f"CROSS-CHECK PASS: {len(rows)} câu → {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
